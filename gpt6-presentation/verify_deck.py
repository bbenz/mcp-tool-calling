"""Reopen, inspect, and locally render the deck; never touches the demo."""

import argparse
import io
import json
import re
from pathlib import Path
from zipfile import ZipFile

import pymupdf
import qrcode
from defusedxml import ElementTree
from PIL import Image, ImageChops, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from content import REPO_URL, SLIDES, TIMING

ROOT = Path(__file__).resolve().parent
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def shapes_in(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from shapes_in(shape.shapes)


def text_frames(shape):
    if shape.has_text_frame:
        yield shape.text_frame
    if shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                yield cell.text_frame


def effective_size(run, paragraph):
    if run.font.size is not None:
        return run.font.size.pt
    if paragraph.font.size is not None:
        return paragraph.font.size.pt
    raise AssertionError("Unexpected master-inherited font; resolve its style before reporting a size.")


def inspect(deck):
    presentation = Presentation(deck)
    assert len(presentation.slides) == len(SLIDES)
    assert abs(presentation.slide_width / 914400 - 13.333333) < 0.001
    assert presentation.slide_height / 914400 == 7.5
    assert sum(segment["end"] - segment["start"] for segment in TIMING) == 1500
    assert TIMING[-1]["start"] == 1380 and TIMING[-1]["end"] == 1500
    assert all(a["end"] == b["start"] for a, b in zip(TIMING, TIMING[1:]))
    summary = []
    fonts = set()
    for number, (slide, data) in enumerate(zip(presentation.slides, SLIDES), 1):
        assert (slide._element.get("show") == "0") == (number > 6)
        assert slide.has_notes_slide
        assert slide.notes_slide.notes_text_frame.text.strip() == data["notes"].strip()
        words = []
        for shape in shapes_in(slide.shapes):
            if shape.shape_type != MSO_SHAPE_TYPE.LINE:
                assert shape.left >= 0 and shape.top >= 0, (number, shape.name)
                assert shape.left + shape.width <= presentation.slide_width, (number, shape.name)
                assert shape.top + shape.height <= presentation.slide_height, (number, shape.name)
            for frame in text_frames(shape):
                words.extend(frame.text.split())
                for paragraph in frame.paragraphs:
                    for run in paragraph.runs:
                        if run.text.strip():
                            size = effective_size(run, paragraph)
                            assert size >= 24, (number, run.text, size)
                            fonts.add(run.font.name or paragraph.font.name)
        summary.append({"slide": number, "title": data["title"], "hidden": number > 6,
                        "projected_words_including_labels": len(words)})
    allowed_urls = {REPO_URL}
    urls = set()
    pictures = [shape for slide in presentation.slides for shape in slide.shapes
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) == 1
    with Image.open(io.BytesIO(pictures[0].image.blob)) as embedded:
        expected = qrcode.make(REPO_URL).convert("RGB")
        assert embedded.size == expected.size
        assert ImageChops.difference(embedded.convert("RGB"), expected).getbbox() is None
    with ZipFile(deck) as archive:
        for name in archive.namelist():
            assert not name.startswith(("ppt/embeddings/", "ppt/activeX/")), name
            if not name.endswith((".xml", ".rels")):
                continue
            xml = archive.read(name).decode("utf-8")
            assert not re.search(r"eyJ[A-Za-z0-9_-]{8,}", xml), name
            assert "-----BEGIN" not in xml, name
            assert not re.search(r"https?://[^\s<>\"']+[?&](?:k|token|key|sig)=", xml, re.I), name
            if name.startswith("ppt/slides/slide") and name.endswith(".xml"):
                root = ElementTree.fromstring(xml)
                assert not any(node.tag.endswith(("}transition", "}timing")) for node in root), name
            if name.endswith(".rels"):
                root = ElementTree.fromstring(xml)
                for relationship in root:
                    if relationship.get("TargetMode") == "External":
                        urls.add(relationship.get("Target"))
    assert urls <= allowed_urls, urls
    return {"structural_checks": "PASS", "slides": summary, "fonts": sorted(fonts),
            "timing_seconds": 1500, "visible": 6, "hidden": len(SLIDES) - 6,
            "hyperlinks": sorted(urls), "qr_matches_public_target": True,
            "render_checks": "NOT RUN"}


def render(deck, report):
    import pywintypes
    import win32com.client

    output = ROOT / "previews"
    output.mkdir(exist_ok=True)
    started = False
    try:
        app = win32com.client.GetActiveObject("PowerPoint.Application")
    except pywintypes.com_error as error:
        if error.hresult != -2147221021:
            raise
        app = win32com.client.DispatchEx("PowerPoint.Application")
        started = True
    document = None
    issues = []
    try:
        document = app.Presentations.Open(str(deck.resolve()), True, False, False)
        for number in range(1, document.Slides.Count + 1):
            slide = document.Slides(number)
            assert bool(slide.SlideShowTransition.Hidden) == (number > 6)
            text_bounds = []
            for index in range(1, slide.Shapes.Count + 1):
                shape = slide.Shapes(index)
                if not shape.HasTextFrame or not shape.TextFrame.HasText:
                    continue
                bound = shape.TextFrame.TextRange
                if (bound.BoundHeight > shape.Height + 3 or bound.BoundWidth > shape.Width + 3):
                    issues.append({"slide": number, "type": "text-overflow",
                                   "text": bound.Text, "shape": [shape.Width, shape.Height],
                                   "text_bounds": [bound.BoundWidth, bound.BoundHeight]})
                rect = (bound.BoundLeft, bound.BoundTop,
                        bound.BoundLeft + bound.BoundWidth, bound.BoundTop + bound.BoundHeight)
                if rect[0] < 35 or rect[1] < 35 or rect[2] > 925 or rect[3] > 505:
                    issues.append({"slide": number, "type": "margin", "text": bound.Text, "bounds": rect})
                for other_rect, other_text in text_bounds:
                    dx = min(rect[2], other_rect[2]) - max(rect[0], other_rect[0])
                    dy = min(rect[3], other_rect[3]) - max(rect[1], other_rect[1])
                    if dx > 2 and dy > 2:
                        issues.append({"slide": number, "type": "text-overlap",
                                       "text": bound.Text, "other": other_text})
                text_bounds.append((rect, bound.Text))
            slide.Export(str(output / f"slide-{number:02}.png"), "PNG", 1920, 1080)
        # Change only the read-only in-memory document so the PDF includes backups.
        for slide in document.Slides:
            slide.SlideShowTransition.Hidden = 0
        document.SaveAs(str(ROOT / "who-can-call-this-mcp-tool.pdf"), 32)
        report["powerpoint_version"] = app.Version
    finally:
        if document is not None:
            document.Saved = True
            document.Close()
        if started and app.Presentations.Count == 0:
            app.Quit()
    with pymupdf.open(ROOT / "who-can-call-this-mcp-tool.pdf") as pdf:
        assert len(pdf) == len(SLIDES)
        for number, page in enumerate(pdf, 1):
            assert page.get_text().strip(), number
            for block in page.get_text("blocks"):
                assert block[0] >= 0 and block[1] >= 0 and block[2] <= page.rect.width + 1, number
        report["pdf_pages_including_backups"] = len(pdf)
    thumbnails = []
    for number in range(1, len(SLIDES) + 1):
        with Image.open(output / f"slide-{number:02}.png") as image:
            assert image.size == (1920, 1080)
            thumb = image.convert("RGB")
            thumb.thumbnail((480, 270))
            thumbnails.append(thumb.copy())
    rows = (len(thumbnails) + 2) // 3
    sheet = Image.new("RGB", (1500, rows * 320 + 20), "#DDE4DE")
    draw = ImageDraw.Draw(sheet)
    font = ImageFont.truetype("C:\\Windows\\Fonts\\calibri.ttf", 22)
    for index, thumb in enumerate(thumbnails):
        x, y = 10 + (index % 3) * 500, 10 + (index // 3) * 320
        sheet.paste(thumb, (x, y))
        draw.text((x + 5, y + 276), f"{index + 1:02} | {'MAIN' if index < 6 else 'BACKUP'}", font=font, fill="#18352F")
    sheet.save(output / "contact-sheet.png")
    report["layout_issues"] = issues
    report["render_checks"] = "PASS" if not issues else "REPAIR REQUIRED"
    return issues


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--render", action="store_true")
    args = parser.parse_args()
    deck = ROOT / "who-can-call-this-mcp-tool.pptx"
    result = inspect(deck)
    issues = render(deck, result) if args.render else []
    (ROOT / "verification-results.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(json.dumps({key: value for key, value in result.items() if key != "slides"}, indent=2))
    if issues:
        raise SystemExit(1)
