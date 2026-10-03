"""Build the offline presentation from its reviewed, local content snapshot."""

import argparse
import io
from pathlib import Path

import qrcode
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from content import REPO_URL, SLIDES
from guide_sections import INTRO

ROOT = Path(__file__).resolve().parent
INK = "18352F"
PAPER = "F7F7F2"
GREEN = "196451"
PALE = "E6EFE9"
WHITE = "FFFFFF"
MUTED = "475B55"
RED = "9A3535"
PINK = "F8E9E5"
GOLD = "F0CB75"
FONT = "Calibri"


def box(slide, x, y, w, h, fill, line=None, radius=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        shape.line.color.rgb = RGBColor.from_string(line)
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def text(slide, value, x, y, w, h, size=28, color=INK, bold=False, font=FONT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for index, line in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = font
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor.from_string(color)
        paragraph.space_after = Pt(5)
    return shape


def arrow(slide, x1, y1, x2, y2, color=GREEN, dashed=False):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    connector.line.color.rgb = RGBColor.from_string(color)
    connector.line.width = Pt(2.5)
    line = connector._element.spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ln")
    if dashed:
        dash = OxmlElement("a:prstDash")
        dash.set("val", "dash")
        line.append(dash)
    head = OxmlElement("a:tailEnd")
    head.set("type", "triangle")
    line.append(head)
    return connector


def panel(slide, title, body, x, y, w, h, strong=False, untrusted=False, size=28):
    fill = INK if strong else PINK if untrusted else PALE
    color = WHITE if strong else RED if untrusted else INK
    box(slide, x, y, w, h, fill, radius=True)
    text(slide, title, x + 0.24, y + 0.22, w - 0.48, 0.6, 28, color, True)
    text(slide, body, x + 0.24, y + 0.96, w - 0.48, h - 1.0, size, color)


def base(presentation, data, number):
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    dark = data["kind"] in {"title", "takeaways"}
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(INK if dark else PAPER)
    color = WHITE if dark else INK
    if data["kind"] != "title":
        text(slide, data["title"], 0.6, 0.52, 12.1, 0.85, 38, color, True)
    label = data.get("label", "MCP DEV SUMMIT TORONTO 2026")
    text(slide, label, 0.6, 6.57, 11.5, 0.4, 24, GOLD if dark else MUTED).name = "Slide footer"
    text(slide, f"{number:02}", 12.08, 6.57, 0.6, 0.4, 24, GOLD if dark else MUTED).name = "Slide number"
    slide.notes_slide.notes_text_frame.text = data["notes"]
    if number > 6:
        slide._element.set("show", "0")
    return slide


def draw_title(slide, data):
    text(slide, "Who Can Call\nThis MCP Tool?", 0.7, 0.8, 8.3, 1.8, 44, WHITE, True)
    text(slide, "OAuth, Resource Binding,\nand Runtime Policy", 0.7, 2.8, 8.2, 1.2, 32, GOLD)
    text(slide, "Brian Benz\nPrincipal AI Advocate, Microsoft", 0.7, 4.25, 8.3, 1.1, 28, WHITE)
    text(slide, "MCP Dev Summit Toronto  |  October 6, 2026", 0.7, 5.55, 11.8, 0.5, 24, WHITE)
    box(slide, 9.1, 1.1, 3.55, 4.1, PALE, radius=True)
    text(slide, "SYNTHETIC\nREFUND", 9.4, 1.45, 2.95, 1.0, 28, GREEN, True)
    text(slide, "User\n+ client\n+ order", 9.4, 2.85, 2.95, 1.8, 32, INK)


def draw_three(slide, data):
    for i, (title, body) in enumerate(data["cards"]):
        panel(slide, title, body, 0.6 + i * 4.1, 1.9, 3.9, 3.8, strong=i == 2)


def draw_architecture(slide, data):
    # The AS branch is separate from the advisory branch; only the API mutates.
    box(slide, 0.6, 2.75, 2.6, 1.7, PINK, RED, True)
    text(slide, "MCP client\nApproval +\nclient model", 0.8, 2.9, 2.2, 1.4, 24, RED, True)
    box(slide, 4.15, 1.6, 3.7, 0.95, PALE, radius=True)
    text(slide, "Authorization server", 4.35, 1.85, 3.3, 0.5, 24, INK, True)
    box(slide, 4.15, 3.0, 3.7, 1.35, INK, radius=True)
    text(slide, "MCP Resource A\nToken checks + policy", 4.38, 3.25, 3.25, 0.9, 24, WHITE, True)
    box(slide, 9.15, 3.0, 3.5, 1.35, PALE, radius=True)
    text(slide, "Upstream API\nIndependent rechecks", 9.35, 3.25, 3.1, 0.9, 24, INK, True)
    box(slide, 9.15, 5.1, 3.5, 0.95, PALE, radius=True)
    text(slide, "Synthetic ledger", 9.35, 5.35, 3.1, 0.5, 24, INK, True)
    box(slide, 4.15, 5.1, 3.7, 0.95, PINK, RED, True)
    text(slide, "Foundry: advice only", 4.35, 5.35, 3.3, 0.5, 24, RED, True)
    arrow(slide, 1.9, 2.7, 1.9, 2.1)
    arrow(slide, 1.9, 2.1, 4.05, 2.1)
    text(slide, "OAuth", 2.25, 1.57, 1.65, 0.42, 24, GREEN)
    arrow(slide, 3.25, 3.6, 4.05, 3.6)
    text(slide, "tools/call", 2.8, 4.48, 2.3, 0.42, 24, GREEN)
    arrow(slide, 6.0, 2.95, 6.0, 2.6)
    text(slide, "Exchange", 6.35, 2.56, 2.25, 0.42, 24, GREEN)
    arrow(slide, 7.9, 3.6, 9.05, 3.6)
    text(slide, "Delegated token", 8.0, 2.48, 3.4, 0.42, 24, GREEN)
    arrow(slide, 10.9, 4.4, 10.9, 5.0)
    text(slide, "Mutation", 11.12, 4.55, 1.6, 0.42, 24, GREEN)
    arrow(slide, 5.95, 4.4, 5.95, 5.0, RED, True)
    text(slide, "Advisory", 6.2, 4.55, 1.6, 0.42, 24, RED)


def draw_map(slide, data, qa=False):
    rows = [
        ("Client", "Human approval", "Not server permission"),
        ("Authorization server", "Identity + audience", "Not object policy"),
        ("MCP server", "Tool + object policy", "Must decide at call time"),
        ("Upstream API", "Mutation checks", "No token passthrough"),
        ("Model", "Advice only", "Never grants authority"),
    ]
    width = 9.0 if qa else 12.1
    heights = 0.73
    columns = (0.23, 3.3, 6.2) if not qa else (0.22, 3.6)
    for i, row in enumerate(rows):
        y = 1.75 + i * (heights + 0.13)
        strong = i == 2
        box(slide, 0.6, y, width, heights, INK if strong else PALE, radius=True)
        color = WHITE if strong else INK
        text(slide, row[0], 0.6 + columns[0], y + 0.17, 3.0, 0.48, 24, color, True)
        text(slide, row[1], 0.6 + columns[1], y + 0.17, 3.9 if qa else 2.9, 0.48, 24, color)
        if not qa:
            text(slide, row[2], 0.6 + columns[2], y + 0.17, 5.3, 0.48, 24, color)
    if qa:
        image = io.BytesIO()
        qrcode.make(REPO_URL).save(image, format="PNG")
        image.seek(0)
        shape = slide.shapes.add_picture(image, Inches(10.05), Inches(1.95), Inches(2.55), Inches(2.55))
        shape._element.nvPicPr.cNvPr.set("descr", f"QR code to public repository: {REPO_URL}")
        text(slide, "Code +\ncontrol map", 10.05, 4.75, 2.55, 1.05, 28, INK, True)
        link = text(slide, "github.com/bbenz/mcp-tool-calling", 0.8, 6.0, 10.3, 0.42, 24, GREEN)
        link.text_frame.paragraphs[0].runs[0].hyperlink.address = REPO_URL
    else:
        text(slide, "Platform filtering is separate from authorization.", 0.8, 6.04, 11.7, 0.45, 24, MUTED)


def draw_takeaways(slide, data):
    for i, (title, body) in enumerate(data["cards"]):
        y = 1.8 + i * 1.45
        box(slide, 0.6, y, 0.7, 0.7, GOLD, radius=True)
        text(slide, str(i + 1), 0.8, y + 0.06, 0.4, 0.5, 28, INK, True)
        text(slide, title, 1.7, y - 0.02, 10.7, 0.6, 32, WHITE, True)
        text(slide, body, 1.7, y + 0.63, 10.7, 0.65, 28, WHITE)


def draw_steps(slide, data):
    for i, (title, body) in enumerate(data["rows"]):
        y = 1.65 + i * 1.1
        box(slide, 0.6, y, 0.65, 0.65, GREEN, radius=True)
        text(slide, str(i + 1), 0.78, y + 0.03, 0.4, 0.5, 28, WHITE, True)
        text(slide, title, 1.65, y, 4.25, 1.03, 28, INK, True)
        text(slide, body, 6.15, y, 6.3, 1.0, 24, INK)


def draw_two(slide, data):
    for i, (title, body) in enumerate(data["cards"]):
        panel(slide, title, body, 0.6 + 6.2 * i, 1.85, 5.9, 3.8,
              strong=data.get("strong_right", False) and i == 1,
              untrusted=data.get("untrusted_left", False) and i == 0,
              size=data.get("size", 28))
    if data.get("bottom"):
        text(slide, data["bottom"], 0.8, 5.94, 11.8, 0.55, 24, GREEN, True)


def draw_policy(slide, data):
    text(slide, "Token validation precedes policy.evaluate.", 0.7, 1.5, 12, 0.55, 28, GREEN, True)
    for i, (title, body) in enumerate(data["rows"]):
        y = 2.3 + i * 0.94
        box(slide, 0.6, y, 12.1, 0.8, PALE, radius=True)
        text(slide, title, 0.85, y + 0.16, 5.5, 0.5, 24, INK, True)
        text(slide, body, 6.4, y + 0.16, 6.0, 0.5, 24, INK)


def draw_refund(slide, data):
    text(slide, "Sam  |  ORD-1001  |  CAD 40.00 = 4000 minor units", 0.7, 1.6, 12, 0.6, 28, INK, True)
    for i, (title, body) in enumerate(data["cards"]):
        panel(slide, title, body, 0.6 + i * 4.1, 2.55, 3.9, 2.25, strong=i == 1, size=28)
    text(slide, "New downstream audience. Same human. No app-only fallback.", 0.8, 5.45, 11.8, 0.8, 28, GREEN, True)


def draw_injection(slide, data):
    rows = [
        ("Untrusted input", "ORD-1005: attacker-written notes"),
        ("Advisory assessment", "Model / filter / offline outcome is reported"),
        ("Forced harness call", "Sam -> ORD-1003, amount_minor=99000"),
        ("Deterministic boundary", "R006 denies; ledger remains unchanged"),
    ]
    draw_steps(slide, {"rows": rows})


def draw_modes(slide, data):
    for i, (title, body) in enumerate(data["cards"]):
        x = 0.6 + (i % 2) * 6.2
        y = 1.7 + (i // 2) * 2.25
        panel(slide, title, body, x, y, 5.9, 2.0, strong=i == 0, size=24)


def draw_audit(slide, data):
    for i, (title, body) in enumerate(data["cards"]):
        panel(slide, title, body, 0.6 + 6.2 * i, 1.7, 5.9, 3.8, size=24)
    text(slide, "Request -> user -> client -> resource -> policy -> result", 0.8, 5.88, 11.8, 0.5, 24, GREEN, True)


def build(output):
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = "Who Can Call This MCP Tool? OAuth, Resource Binding, and Runtime Policy"
    presentation.core_properties.author = "Brian Benz"
    presentation.core_properties.subject = "MCP Dev Summit Toronto 2026 - demo-first control map"
    presentation.core_properties.keywords = "MCP, OAuth, runtime authorization, synthetic refund demo"
    presentation.core_properties.comments = "Generated from reviewed local sources; no live demo executed."
    renderers = {
        "title": draw_title, "three": draw_three, "architecture": draw_architecture,
        "map": draw_map, "takeaways": draw_takeaways, "qa": lambda s, d: draw_map(s, d, True),
        "steps": draw_steps, "two": draw_two, "policy": draw_policy,
        "refund": draw_refund, "injection": draw_injection, "modes": draw_modes,
        "audit": draw_audit,
    }
    for number, data in enumerate(SLIDES, 1):
        slide = base(presentation, data, number)
        renderers[data["kind"]](slide, data)
        for shape in list(slide.shapes):
            if shape.name in {"Slide footer", "Slide number"}:
                slide.shapes._spTree.remove(shape._element)
                slide.shapes._spTree.append(shape._element)
    presentation.save(output)
    guide = [INTRO, "\n| Slide | Title / purpose | Visibility |\n|---:|---|---|\n"]
    for number, data in enumerate(SLIDES, 1):
        guide.append(f"| {number} | {data['title']} | {'Main' if number <= 6 else 'Hidden backup'} |\n")
    for number, (slide, data) in enumerate(zip(presentation.slides, SLIDES), 1):
        guide.append(f"\n### Slide {number}: {data['title']}\n\n")
        guide.append("**Projected text (including evidence labels):**\n\n")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip() and shape.name != "Slide number":
                guide.append("- " + shape.text.replace("\n", " / ") + "\n")
        guide.append("\n**Saved speaker notes:**\n\n" + data["notes"].replace("\n", "\n\n") + "\n")
    (output.parent / "presenter-guide.md").write_text("".join(guide), encoding="utf-8")
    print(f"Created {output} ({len(SLIDES)} slides, 6 visible)")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=ROOT / "who-can-call-this-mcp-tool.pptx")
    args = parser.parse_args()
    build(args.output.resolve())
